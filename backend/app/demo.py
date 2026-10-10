"""Generate synthetic evidence workbooks; no private sources or credentials read."""

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook


def create_demo(directory: Path) -> list[Path]:
    fixtures: dict[str, dict[str, list[list[Any]]]] = {
        "synthetic-feedback.xlsx": {
            "Feedback_Explorer": [
                [
                    "ID",
                    "Source",
                    "Brand",
                    "Sales Description",
                    "Country",
                    "Derivate (e-code)",
                    "Engine Type",
                    "Feedback Type",
                    "Customer Feedback",
                    "Vfc level2 3 Name",
                    "Vfc level2 Name",
                    "Vfc level3 Name",
                ],
                [
                    "synthetic-1",
                    "synthetic_survey",
                    "Synthetic",
                    None,
                    "Germany",
                    "G60",
                    "unknown-code",
                    None,
                    "Synthetic: the seat is uncomfortable on long journeys",
                    None,
                    "comfort",
                    None,
                ],
                [
                    "synthetic-2",
                    "synthetic_survey",
                    "Synthetic",
                    None,
                    "Germany",
                    "G60",
                    "unknown-code",
                    None,
                    "Synthetic: the seat is comfortable on short journeys",
                    None,
                    "comfort",
                    None,
                ],
            ]
        },
        "synthetic-study.xlsx": {
            "US_2025": [
                ["Synthetic comfort", None, "Model G60"],
                ["Sample total", "Synthetic comfort", 10],
                ["Satisfied", "Synthetic comfort", 0.6],
                ["Dissatisfied", "Synthetic comfort", 0.4],
            ]
        },
        "synthetic-sales.xlsx": {
            "G60_G68": [
                ["market", "market_code", "Sales 2025", "Sales 2030"],
                ["USA", "US", 0, 100],
            ]
        },
    }
    root = directory.resolve()
    targets = [root / filename for filename in fixtures]
    if any(path.exists() for path in targets):
        raise FileExistsError(
            "Synthetic fixture already exists; choose an empty destination"
        )
    root.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for filename, sheets in fixtures.items():
        workbook = Workbook()
        del workbook[workbook.sheetnames[0]]
        for title, rows in sheets.items():
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(row)
        path = root / filename
        with path.open("xb") as stream:
            workbook.save(stream)
        workbook.close()
        written.append(path)
    return written


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate synthetic feedback, study and sales workbooks"
    )
    parser.add_argument("--directory", type=Path, required=True)
    args = parser.parse_args()
    paths = create_demo(args.directory)
    print(json.dumps({"synthetic": True, "files": [str(path) for path in paths]}))


if __name__ == "__main__":
    main()
