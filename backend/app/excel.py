"""Normalize workbook facts without executing formulas or inventing identities."""

import math
import re
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ingest_types import (
    ContextMetric,
    CoverageIssue,
    FeedbackAnnotation,
    FeedbackRecord,
    ParseResult,
)
from app.wire import CellLocator, Scope

REVISION = "excel-v2"
HEADERS = {
    "id": "id",
    "source": "source",
    "country": "country",
    "derivate (e-code)": "vehicle",
    "derivative (e-code)": "vehicle",
    "customer feedback": "text",
}
COUNTRIES = {
    "germany": "DE",
    "de": "DE",
    "deutschland": "DE",
    "france": "FR",
    "fr": "FR",
    "italy": "IT",
    "it": "IT",
    "spain": "ES",
    "es": "ES",
    "netherlands": "NL",
    "nl": "NL",
    "austria": "AT",
    "at": "AT",
    "belgium": "BE",
    "be": "BE",
    "poland": "PL",
    "pl": "PL",
    "sweden": "SE",
    "se": "SE",
    "china": "CN",
    "cn": "CN",
    "usa": "US",
    "us": "US",
    "united states": "US",
    "united states of america": "US",
    "europe": "EU",
    "eu": "EU",
    "united kingdom": "GB",
    "uk": "GB",
    "gb": "GB",
    "japan": "JP",
    "jp": "JP",
    "south korea": "KR",
}
EU_MARKETS = {"DE", "FR", "IT", "ES", "NL", "AT", "BE", "PL", "SE"}


def text(value: Any) -> str:
    return "" if value is None else str(value)


def locator(sheet: str, row: int, columns: list[int]) -> CellLocator:
    return CellLocator(
        sheet=sheet, row=row, cells=[f"{get_column_letter(c)}{row}" for c in columns]
    )


def issue(result: ParseResult, code: str, location: CellLocator | None = None) -> None:
    result.coverage.append(
        CoverageIssue(code=code, message=code.replace("_", " "), locator=location)
    )
    if result.status == "ready":
        result.status = "partial"


def scope_for(
    country: Any, vehicle: Any, result: ParseResult, location: CellLocator
) -> Scope:
    country_text = text(country).strip()
    market = COUNTRIES.get(country_text.casefold())
    if country_text and market is None:
        issue(result, "unknown_country_mapping", location)
    vehicles = re.findall(r"\b[A-Z]\d{2}\b", text(vehicle).upper())
    if text(vehicle).strip() and not vehicles:
        issue(result, "unknown_vehicle_mapping", location)
    cohorts = [market] if market else []
    if market in EU_MARKETS:
        cohorts.append("EU")
    return Scope(
        vehicle_ids=sorted(set(vehicles)),
        markets=[market] if market else [],
        cohort_ids=cohorts,
    )


def feedback(
    sheet: Any,
    raw_sheet: Any,
    result: ParseResult,
    namespace: str,
    aliases: dict[str, str],
) -> None:
    rows = sheet.iter_rows()
    raw_rows = raw_sheet.iter_rows()
    headers = [text(cell.value) for cell in next(rows, ())]
    next(raw_rows, ())
    mapped = [aliases.get(" ".join(header.casefold().split())) for header in headers]
    required = {"id", "source", "country", "vehicle", "text"}
    if not required.issubset(mapped) or any(
        mapped.count(name) != 1 for name in required
    ):
        result.status = "unsupported"
        issue(
            result,
            "mapping_required",
            locator(sheet.title, 1, list(range(1, len(headers) + 1))),
        )
        return
    columns = {name: mapped.index(name) for name in required}
    records: dict[tuple[str, str], FeedbackRecord] = {}
    for row_number, (row, raw_row) in enumerate(zip(rows, raw_rows, strict=True), 2):
        values = [cell.value for cell in row]
        if not any(value is not None for value in values):
            continue
        location = locator(sheet.title, row_number, list(range(1, len(headers) + 1)))
        scope = scope_for(
            values[columns["country"]], values[columns["vehicle"]], result, location
        )
        missing_cache = any(
            cell.data_type == "f" and values[index] is None
            for index, cell in enumerate(raw_row)
        )
        if missing_cache:
            issue(result, "missing_formula_cache", location)
        annotation = FeedbackAnnotation(
            text=text(values[columns["text"]]),
            locator=location,
            scope=scope,
            classifications={
                header: text(raw_row[index].value)
                if raw_row[index].value is not None
                else None
                for index, header in enumerate(headers)
            },
        )
        external_id = text(values[columns["id"]]).strip()
        source = text(values[columns["source"]]).strip()
        if not external_id or not source or len(external_id) > 200 or len(source) > 200:
            issue(result, "missing_record_identity", location)
            result.unresolved_annotations.append(annotation)
            continue
        if not annotation.text.strip():
            issue(result, "missing_feedback_text", location)
            result.unresolved_annotations.append(annotation)
            continue
        key = (source, external_id)
        record = records.get(key)
        if record is None:
            record = FeedbackRecord(
                dataset_namespace=namespace,
                source_namespace=source,
                external_record_id=external_id,
                text=annotation.text,
                annotations=[],
                scope=scope.model_copy(deep=True),
                text_variants=[annotation.text],
            )
            records[key] = record
        if annotation.text not in record.text_variants:
            record.text_variants.append(annotation.text)
            record.limitations = ["conflicting_feedback_text"]
            issue(result, "conflicting_feedback_text", location)
        record.annotations.append(annotation)
        for name in ("vehicle_ids", "markets", "cohort_ids"):
            setattr(
                record.scope,
                name,
                sorted(set(getattr(record.scope, name)) | set(getattr(scope, name))),
            )
    result.records.extend(records.values())


def parse_excel(
    data: bytes,
    *,
    dataset_namespace: str,
    reference_year: int = 2026,
    header_aliases: dict[str, str] | None = None,
) -> ParseResult:
    result = ParseResult(
        kind="unknown",
        parser_revision=REVISION,
        parser_versions={"openpyxl": openpyxl.__version__},
    )
    book = load_workbook(
        BytesIO(data), read_only=True, data_only=True, keep_links=False
    )
    raw_book = load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    try:
        if "Feedback_Explorer" in book.sheetnames:
            result.kind = "feedback"
            aliases = HEADERS | {
                " ".join(key.casefold().split()): value
                for key, value in (header_aliases or {}).items()
            }
            feedback(
                book["Feedback_Explorer"],
                raw_book["Feedback_Explorer"],
                result,
                dataset_namespace,
                aliases,
            )
        elif any(name.startswith(("CN_EU", "US_")) for name in book.sheetnames):
            result.kind = "study"
            for sheet in book:
                if sheet.title.startswith(("CN_EU", "US_")):
                    study(sheet, raw_book[sheet.title], result)
                else:
                    issue(result, "mapping_required", locator(sheet.title, 1, [1]))
            if not result.metrics:
                result.status = "unsupported"
        else:
            result.status = "unsupported"
            issue(result, "mapping_required")
        return result
    finally:
        book.close()
        raw_book.close()


def number(value: Any, result: ParseResult, location: CellLocator) -> float | None:
    if value is None:
        issue(result, "missing_metric_value", location)
        return None
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
    ):
        issue(result, "invalid_numeric_value", location)
        return None
    return float(value)


def study_metric(
    result: ParseResult,
    sheet: str,
    rows: list[list[Any]],
    references: list[tuple[int, int]],
    topic: str,
    metric: str,
    value: float | None,
    scope: Scope,
    proportion_base: float | None = None,
    limitations: list[str] | None = None,
) -> ContextMetric:
    cells = [f"{get_column_letter(column)}{row}" for row, column in references]
    raw_cells = {}
    for address, (row, column) in zip(cells, references, strict=True):
        raw = rows[row - 1][column - 1]
        raw_cells[address] = (
            raw if raw is None or isinstance(raw, (str, int, float, bool)) else str(raw)
        )
    item = ContextMetric(
        source_kind="study",
        topic=topic,
        metric=metric,
        value=value,
        unit="fraction" if metric == "fraction" else None,
        denominator=None,
        proportion_base=proportion_base,
        raw_cells=raw_cells,
        locator=CellLocator(sheet=sheet, row=references[0][0], cells=cells),
        scope=scope,
        limitations=[
            "unknown_sample_size",
            *([] if metric == "fraction" else ["unknown_mean_scale"]),
            *(limitations or []),
        ],
    )
    result.metrics.append(item)
    return item


def study(sheet: Any, raw_sheet: Any, result: ParseResult) -> None:
    rows = [list(row) for row in sheet.values]
    raw_rows = [list(row) for row in raw_sheet.values]
    year_match = re.search(r"\b(20\d{2})\b", sheet.title.replace("_", " "))
    year = int(year_match[1]) if year_match else None
    if year is None:
        issue(result, "unknown_study_year", locator(sheet.title, 1, [1]))
    if sheet.title.startswith("CN_EU"):
        mean_study(sheet.title, rows, raw_rows, year, result)
    else:
        fraction_study(sheet.title, rows, raw_rows, year, result)


def mean_study(
    sheet: str,
    rows: list[list[Any]],
    raw_rows: list[list[Any]],
    year: int | None,
    result: ParseResult,
) -> None:
    header_row = next(
        (
            index
            for index, row in enumerate(rows)
            if any(re.search(r"\b[A-Z]\d{2}\b", text(value)) for value in row[1:])
        ),
        None,
    )
    if header_row is None or header_row == 0:
        issue(result, "mapping_required", locator(sheet, 1, [1]))
        return
    country = None
    country_column = 0
    columns: dict[int, tuple[Scope, int]] = {}
    for column, value in enumerate(rows[header_row][1:], 1):
        declared_country = rows[header_row - 1][column]
        if declared_country is not None:
            country, country_column = declared_country, column
        if not text(value).strip():
            continue
        location = locator(sheet, header_row + 1, [column + 1])
        scope = scope_for(country, value, result, location)
        scope.year = year
        if scope.markets and scope.vehicle_ids:
            columns[column] = (scope, country_column)
        else:
            issue(result, "mapping_required", location)
    topic = ""
    topic_row = header_row + 1
    for index in range(header_row, len(rows)):
        row = rows[index]
        label = text(row[0]).strip()
        if label and label.casefold() != "mean":
            topic, topic_row = label, index + 1
            continue
        if label.casefold() != "mean":
            continue
        if not topic:
            issue(result, "mapping_required", locator(sheet, index + 1, [1]))
            continue
        for column, (scope, market_column) in columns.items():
            location = locator(sheet, index + 1, [column + 1])
            value = number(row[column], result, location)
            references = [
                (index + 1, column + 1),
                (topic_row, 1),
                (header_row + 1, column + 1),
                (header_row, market_column + 1),
            ]
            study_metric(
                result, sheet, raw_rows, references, topic, "mean", value, scope
            )


def fraction_study(
    sheet: str,
    rows: list[list[Any]],
    raw_rows: list[list[Any]],
    year: int | None,
    result: ParseResult,
) -> None:
    if not rows:
        issue(result, "mapping_required", locator(sheet, 1, [1]))
        return
    columns: dict[int, Scope] = {}
    for column, value in enumerate(rows[0][2:], 2):
        if not text(value).strip():
            continue
        scope = scope_for("US", value, result, locator(sheet, 1, [column + 1]))
        scope.year = year
        if scope.vehicle_ids:
            columns[column] = scope
    if not columns:
        issue(result, "mapping_required", locator(sheet, 1, [1]))
        return
    topic = ""
    topic_row = 1
    totals: dict[int, tuple[float | None, int]] = {}
    groups: dict[int, list[ContextMetric]] = {}

    def check_totals() -> None:
        # Published fractions tolerate two percentage points of rounding, never renormalize.
        for items in groups.values():
            if (
                any(item.value is None for item in items)
                or abs(sum(item.value or 0 for item in items) - 1) > 0.02
            ):
                issue(result, "fraction_total_anomaly", items[0].locator)
                for item in items:
                    item.limitations.append("fraction_total_anomaly")
        groups.clear()

    for index, row in enumerate(rows):
        label = text(row[0]).strip()
        if not any(text(value).strip() for value in row):
            continue
        if label.casefold() == "sample total":
            totals.clear()
            for column in columns:
                base = number(
                    row[column], result, locator(sheet, index + 1, [column + 1])
                )
                if base is not None and base <= 0:
                    issue(
                        result,
                        "invalid_proportion_base",
                        locator(sheet, index + 1, [column + 1]),
                    )
                    base = None
                totals[column] = (base, index + 1)
            continue
        if label and (index == 0 or not any(value is not None for value in row[1:])):
            check_totals()
            topic, topic_row = label, index + 1
            totals.clear()
            continue
        if not label or not topic or not totals:
            issue(result, "mapping_required", locator(sheet, index + 1, [1]))
            continue
        for column, scope in columns.items():
            location = locator(sheet, index + 1, [column + 1])
            value = number(row[column], result, location)
            limitations = []
            if value is not None and not 0 <= value <= 1:
                issue(result, "fraction_out_of_range", location)
                limitations.append("fraction_out_of_range")
                value = None
            base, base_row = totals[column]
            references = [
                (index + 1, column + 1),
                (index + 1, 1),
                (topic_row, 1),
                (1, column + 1),
                (base_row, column + 1),
            ]
            item = study_metric(
                result,
                sheet,
                raw_rows,
                references,
                topic,
                "fraction",
                value,
                scope,
                base,
                limitations,
            )
            groups.setdefault(column, []).append(item)
    check_totals()
