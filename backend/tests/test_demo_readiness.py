import asyncio

import pytest
from openpyxl import load_workbook

from app.demo import create_demo


def test_demo_creates_only_synthetic_fixtures_and_refuses_overwrite(tmp_path):
    paths = create_demo(tmp_path)
    assert {path.name for path in paths} == {
        "synthetic-feedback.xlsx",
        "synthetic-study.xlsx",
        "synthetic-sales.xlsx",
    }
    feedback = load_workbook(tmp_path / "synthetic-feedback.xlsx", read_only=True)
    assert feedback["Feedback_Explorer"].max_row == 3
    assert feedback["Feedback_Explorer"]["I2"].value.startswith("Synthetic:")
    feedback.close()
    sales = load_workbook(tmp_path / "synthetic-sales.xlsx", read_only=True)
    assert sales["G60_G68"]["C2"].value == 0
    sales.close()
    before = {path.name: path.read_bytes() for path in paths}
    with pytest.raises(FileExistsError):
        create_demo(tmp_path)
    assert before == {path.name: path.read_bytes() for path in paths}


@pytest.mark.parametrize("healthy,expected", [(True, 200), (False, 503)])
def test_readiness_reports_dependencies_and_config_flags_without_secret_values(
    monkeypatch, healthy, expected
):
    import app.readiness as readiness
    from app.config import Settings

    settings = Settings(
        _env_file=None,
        openai_api_key="synthetic-private-key",
        renderer_url="https://renderer.example.invalid",
        renderer_token="synthetic-private-token",
    )
    monkeypatch.setattr(readiness, "Settings", lambda: settings)
    monkeypatch.setattr(readiness, "database_ready", lambda settings: healthy)

    async def temporal(settings):
        return healthy

    monkeypatch.setattr(readiness, "temporal_ready", temporal)
    payload, status = asyncio.run(readiness.check_readiness())
    assert status == expected
    assert payload["dependencies"] == {"database": healthy, "temporal": healthy}
    assert payload["configured_capabilities"]["renderer"]
    assert "synthetic-private" not in str(payload)
